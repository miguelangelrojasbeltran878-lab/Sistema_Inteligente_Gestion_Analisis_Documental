import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUTPUT_DIR = r"c:\Users\Lenovo\OneDrive\Documentos\Sistema_Inteligente_Gestion_Analisis_Documental\03_DESARROLLO"

def create_base_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
    return doc

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_block(doc, title, subtitle):
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(2)
    r1 = title_p.add_run(title)
    r1.font.name = 'Arial'
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(14, 56, 122)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(12)
    r2 = sub_p.add_run(subtitle)
    r2.font.name = 'Arial'
    r2.font.size = Pt(10)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(90, 90, 90)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(14, 56, 122)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 80, 160)
    return p

def add_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.bold = True
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = RGBColor(20, 20, 20)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.bold = True
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = RGBColor(20, 20, 20)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    r = p.add_run(code_text)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(30, 30, 30)

def add_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "0E387A")
        set_cell_margins(hdr_cells[i], 100, 100, 120, 120)
        for p in hdr_cells[i].paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F4F6F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], 80, 80, 100, 100)
            for p in row_cells[c_idx].paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(40, 40, 40)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

# ==============================================================================
# DOC 1: CONFIGURACIÓN Y ESTRUCTURA
# ==============================================================================
def build_doc1():
    doc = create_base_doc()
    add_header_block(doc, "GUÍA DE CONFIGURACIÓN DEL ENTORNO, ESTRUCTURA Y GITFLOW",
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nVersión 1.0.0 - Línea Base de Desarrollo")

    add_heading_1(doc, "1. Descripción del Entorno de Desarrollo")
    add_heading_2(doc, "1.1 Tecnologías y Versiones Oficiales")
    headers = ["Herramienta / Entorno", "Versión Oficial", "Propósito en el Stack"]
    rows = [
        ["Node.js", "v20.17.0 LTS", "Entorno para frontend React y backend de servicios"],
        ["Python", "3.11.9+", "Runtime para IA, OCR Tesseract, Sentence-Transformers y RAG"],
        ["PostgreSQL", "16.3", "Base de datos relacional principal"],
        ["pgvector", "v0.7.0+", "Extensión PostgreSQL para vectores y búsqueda HNSW"],
        ["Tesseract OCR", "5.3.4+", "Motor OCR para extracción de texto en imágenes escaneadas"],
        ["Docker & Compose", "26.1.0+ / v2.27+", "Contenerización y orquestación multi-contenedor"],
        ["Git", "2.45.0+", "Control de versiones distribuido"]
    ]
    add_table(doc, headers, rows, [1.5, 1.5, 3.8])

    add_heading_2(doc, "1.2 Plantilla de Variables de Entorno (.env.example)")
    env_code = """NODE_ENV=development
APP_PORT=8080
API_PREFIX=/api/v1
JWT_SECRET_KEY=uts_super_secret_jwt_key_2026_secure_sha256_minimum_64_chars!
JWT_EXPIRATION_SECONDS=3600
BCRYPT_SALT_ROUNDS=12
DB_HOST=localhost
DB_PORT=5432
DB_NAME=uts_documental_db
DB_USER=postgres
DB_PASSWORD=uts_postgres_secure_2026
AI_SERVICE_URL=http://localhost:8000
EMBEDDING_MODEL_NAME=sentence-transformers/all-MiniLM-L6-v2
EMBEDDING_DIMENSION=384
GEMINI_API_KEY=AIzaSyD_EXAMPLE_KEY_UTS_DOCUMENTAL_2026
UPLOAD_DIR=./uploads/documents
MAX_FILE_SIZE_BYTES=20971520"""
    add_code_block(doc, env_code)

    add_heading_2(doc, "1.3 Políticas de Seguridad y .gitignore")
    add_bullet(doc, "Exclusión estricta de .env, node_modules/, __pycache__/, .venv/ y la carpeta uploads/.", "Reglas de Exclusión: ")

    add_heading_1(doc, "2. Estructura Completa del Código Fuente")
    add_bullet(doc, "/frontend: Aplicación web React 18, Vite, TypeScript, TailwindCSS y Lucide Icons.", "Módulo Frontend: ")
    add_bullet(doc, "/backend: Capa REST API (Clean Architecture), JWT, RBAC y orquestación.", "Módulo Backend: ")
    add_bullet(doc, "/ai: Microservicio Python FastAPI con Sentence-Transformers, LLM y motor RAG.", "Módulo IA: ")
    add_bullet(doc, "/document_processing: Extractores PyMuPDF, python-docx, OCR Tesseract y chunker.", "Módulo Ingestión: ")
    add_bullet(doc, "/database: Migraciones SQL, extensión pgvector y semillas de prueba.", "Módulo Base de Datos: ")
    add_bullet(doc, "/tests: Suite de pruebas unitarias, integración y Playwright E2E.", "Módulo Pruebas: ")

    add_heading_1(doc, "3. Control de Versiones y Estrategia Git")
    add_p(doc, "Se utiliza Gitflow con ramas main (producción v1.0.0), develop (integración), feature/* (desarrollos de módulos) y fix/* (correcciones). Commits bajo estándar Conventional Commits (feat, fix, docs, refactor, test, chore).")

    add_heading_1(doc, "4. Bitácora de Desarrollo y Registro de Avances")
    h_headers = ["Hito", "Fechas", "Componente", "Responsable", "Entregable Completado"]
    h_rows = [
        ["H-01", "10/08 - 15/08", "Arquitectura & Entorno", "Tech Lead", "Monorepo, Docker Compose, .env y DB inicial."],
        ["H-02", "16/08 - 22/08", "Seguridad & CRUD", "Dev Backend", "Auth JWT, BCrypt, carpetas y carga de archivos 20MB."],
        ["H-03", "23/08 - 29/08", "Extracción & OCR", "Dev IA/NLP", "Parsers PDF/DOCX/TXT, OCR Tesseract y chunking."],
        ["H-04", "30/08 - 05/09", "Inferencia IA & RAG", "Dev IA/NLP", "Clasificador 3 clases, resumen, JSON y motor RAG."],
        ["H-05", "06/09 - 12/09", "Frontend SPA", "Dev Frontend", "Vistas React (Login, Dashboard, Detalle, Chat RAG)."],
        ["H-06", "13/09 - 18/09", "Integración & QA", "Lead QA", "Pruebas E2E, latencias < 2.5s y auditoría final."]
    ]
    add_table(doc, h_headers, h_rows, [0.7, 1.2, 1.5, 1.2, 2.2])

    doc.save(os.path.join(OUTPUT_DIR, "01_Configuracion_y_Estructura.docx"))
    print("Dev Doc 1 generated.")

# ==============================================================================
# DOC 2: DESARROLLO BACKEND
# ==============================================================================
def build_doc2():
    doc = create_base_doc()
    add_header_block(doc, "DESARROLLO DE LA CAPA BACKEND Y API REST",
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nArquitectura: Clean Architecture en Capas Desacopladas")

    add_heading_1(doc, "1. Arquitectura del Backend y Patrón de Diseño")
    add_p(doc, "El backend implementa el patrón Controller-Service-Repository con middlewares de autenticación JWT y guardias RBAC. La capa de controladores se encarga del transporte HTTP y códigos de estado; los servicios orquestan la lógica de negocio y las llamadas asíncronas al microservicio de IA; los repositorios realizan consultas SQL parametrizadas a PostgreSQL.")

    add_heading_1(doc, "2. Implementación de Módulos Nucleares")
    
    add_heading_2(doc, "2.1 Middleware de Autenticación JWT y Guardias RBAC")
    auth_code = """export const authenticateJWT = (req: Request, res: Response, next: NextFunction) => {
  const authHeader = req.headers.authorization;
  if (!authHeader || !authHeader.startsWith('Bearer ')) {
    return res.status(401).json({ success: false, error: { code: 'UNAUTHORIZED' } });
  }
  const token = authHeader.split(' ')[1];
  try {
    req.user = jwt.verify(token, process.env.JWT_SECRET_KEY!) as AuthenticatedUser;
    next();
  } catch (err) {
    return res.status(401).json({ success: false, error: { code: 'INVALID_TOKEN' } });
  }
};

export const requireRoles = (roles: string[]) => {
  return (req: Request, res: Response, next: NextFunction) => {
    if (!req.user || !roles.includes(req.user.role)) {
      return res.status(403).json({ success: false, error: { code: 'FORBIDDEN' } });
    }
    next();
  };
};"""
    add_code_block(doc, auth_code)

    add_heading_2(doc, "2.2 Middleware de Carga y Validación de Archivos (Multer)")
    multer_code = """export const uploadDocumentMiddleware = multer({
  storage: multer.diskStorage({
    destination: (req, file, cb) => cb(null, './uploads/documents'),
    filename: (req, file, cb) => cb(null, `${crypto.randomUUID()}${path.extname(file.originalname)}`)
  }),
  limits: { fileSize: 20 * 1024 * 1024 }, // 20 MB
  fileFilter: (req, file, cb) => {
    const ext = path.extname(file.originalname).toLowerCase();
    ['.pdf', '.docx', '.txt'].includes(ext) ? cb(null, true) : cb(new Error('INVALID_FILE_TYPE'));
  }
}).single('file');"""
    add_code_block(doc, multer_code)

    add_heading_2(doc, "2.3 Invocador del Microservicio de IA")
    ai_call_code = """export class AIClientService {
  async triggerDocumentProcessing(documentId: string, filePath: string) {
    const res = await axios.post(`${process.env.AI_SERVICE_URL}/internal/v1/process-document`, 
      { document_id: documentId, file_path: filePath },
      { headers: { 'X-Internal-Token': process.env.AI_INTERNAL_API_KEY }, timeout: 60000 }
    );
    return res.data;
  }
}"""
    add_code_block(doc, ai_call_code)

    add_heading_2(doc, "2.4 Módulo de Dashboard y Métricas")
    dash_code = """export class DashboardService {
  async getMetrics() {
    const statusRes = await pool.query(`SELECT status, COUNT(*)::int FROM documents GROUP BY status;`);
    const catRes = await pool.query(`SELECT category, COUNT(*)::int FROM document_analysis GROUP BY category;`);
    const totalRes = await pool.query(`SELECT COUNT(*)::int as total, COALESCE(SUM(file_size), 0)::bigint as bytes FROM documents;`);
    return { status_distribution: statusRes.rows, category_distribution: catRes.rows, totals: totalRes.rows[0] };
  }
}"""
    add_code_block(doc, dash_code)

    add_heading_1(doc, "3. Gestión Global de Excepciones y Logging de Auditoría")
    add_p(doc, "El middleware de errores captura excepciones (INVALID_FILE_TYPE -> 400, LIMIT_FILE_SIZE -> 413, INVALID_CREDENTIALS -> 401, 500 para excepciones no controladas) y registra automáticamente las trazas de fallo en la tabla processing_logs de PostgreSQL.")

    doc.save(os.path.join(OUTPUT_DIR, "02_Desarrollo_Backend.docx"))
    print("Dev Doc 2 generated.")

# ==============================================================================
# DOC 3: DESARROLLO FRONTEND
# ==============================================================================
def build_doc3():
    doc = create_base_doc()
    add_header_block(doc, "DESARROLLO DE LA CAPA CLIENTE FRONTEND (SPA)",
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nStack: React 18 + TypeScript + Vite + Tailwind CSS")

    add_heading_1(doc, "1. Arquitectura de Cliente y Organización Modular")
    add_p(doc, "Estructura modular dividida en: /components (common, dashboard, documents, analysis, chat), /context (AuthContext), /hooks (useAuth, useDocuments, useRAGChat), /pages (Login, Dashboard, Explorer, Detail, Chat), /services (Axios API clients) y /types.")

    add_heading_1(doc, "2. Implementación de Vistas y Experiencia de Usuario")
    
    add_heading_2(doc, "2.1 AuthContext y Rutas Protegidas")
    auth_ctx_code = """export const AuthProvider: React.FC<{ children: React.ReactNode }> = ({ children }) => {
  const [token, setToken] = useState<string | null>(localStorage.getItem('uts_access_token'));
  const [user, setUser] = useState<User | null>(null);

  const login = async (email: string, pass: string) => {
    const res = await api.post('/auth/login', { email, password: pass });
    setToken(res.data.data.access_token);
    setUser(res.data.data.user);
    localStorage.setItem('uts_access_token', res.data.data.access_token);
  };
  return <AuthContext.Provider value={{ user, token, login, isAuthenticated: !!token }}>{children}</AuthContext.Provider>;
};"""
    add_code_block(doc, auth_ctx_code)

    add_heading_2(doc, "2.2 Componente FileDropzone (Drag & Drop)")
    drop_code = """export const FileDropzone = ({ folderId, onUploadSuccess }) => {
  const validateAndUpload = async (file: File) => {
    if (!['pdf','docx','txt'].includes(file.name.split('.').pop()?.toLowerCase() || '')) return alert('Formato no permitido');
    if (file.size > 20 * 1024 * 1024) return alert('Supera 20 MB');
    const fd = new FormData();
    fd.append('file', file);
    fd.append('folder_id', folderId);
    await api.post('/documents/upload', fd, { headers: { 'Content-Type': 'multipart/form-data' } });
    onUploadSuccess();
  };
  return <div onDrop={(e) => { e.preventDefault(); validateAndUpload(e.dataTransfer.files[0]); }}>Arrastra tu archivo aquí</div>;
};"""
    add_code_block(doc, drop_code)

    add_heading_2(doc, "2.3 Detalle Documental con Resumen y JSON de Entidades")
    add_p(doc, "Despliega el resumen ejecutivo en formato Markdown enriquecido, tarjeta de badges de categoría (Administrativo, Financiero, Técnico/Legal) y visor de pares clave-valor extraídos.")

    add_heading_2(doc, "2.4 Módulo de Chat RAG con Citas de Sustento")
    rag_ui_code = """export const RAGChatWindow = () => {
  const [messages, setMessages] = useState([]);
  const [input, setInput] = useState('');
  const handleSend = async () => {
    const res = await api.post('/documents/query', { question: input });
    setMessages(prev => [...prev, { text: input, sender: 'user' }, { text: res.data.data.answer, sources: res.data.data.sources, sender: 'ai' }]);
  };
  return <div>{/* Historial de mensajes con tarjetas de fuentes [Archivo, Página] */}</div>;
};"""
    add_code_block(doc, rag_ui_code)

    add_heading_1(doc, "3. Manejo de Estados Globales y Notificaciones")
    add_p(doc, "Integración de componentes ToastAlert para feedback de éxito/error y Spinners/Skeletons para estados asíncronos durante la inferencia de IA.")

    doc.save(os.path.join(OUTPUT_DIR, "03_Desarrollo_Frontend.docx"))
    print("Dev Doc 3 generated.")

# ==============================================================================
# DOC 4: IA Y PROCESAMIENTO DOCUMENTAL
# ==============================================================================
def build_doc4():
    doc = create_base_doc()
    add_header_block(doc, "INGENIERÍA DE INTELIGENCIA ARTIFICIAL, NLP Y MOTOR RAG",
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nStack: Python 3.11 + FastAPI + PyMuPDF + Tesseract + Sentence-Transformers + pgvector")

    add_heading_1(doc, "1. Pipeline Técnico de Ingestión Documental")
    add_p(doc, "Flujo: Ingesta binaria -> Extracción (PyMuPDF para PDF nativo, python-docx para DOCX, OCR Tesseract para PDFs escaneados) -> Normalización -> Chunking (500 tokens, 100 overlap).")

    add_heading_2(doc, "1.1 Código del Extractor Multiformato con OCR Fallback")
    ext_code = """class DocumentExtractor:
    @staticmethod
    def extract_text(file_path: str) -> dict:
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".txt":
            with open(file_path, "r", encoding="utf-8") as f: return {"full_text": f.read()}
        elif ext == ".docx":
            doc = docx.Document(file_path)
            return {"full_text": "\\n".join([p.text for p in doc.paragraphs if p.text.strip()])}
        elif ext == ".pdf":
            doc = fitz.open(file_path)
            text_list = []
            for page in doc:
                t = page.get_text().strip()
                if len(t) < 40: # Escaneo -> OCR Tesseract
                    pix = page.get_pixmap(dpi=200)
                    t = pytesseract.image_to_string(Image.open(io.BytesIO(pix.tobytes("png"))), lang="spa")
                text_list.append(t)
            return {"full_text": "\\n\\n".join(text_list)}"""
    add_code_block(doc, ext_code)

    add_heading_1(doc, "2. Integración con Modelos de IA (Clasificación, Resumen y JSON)")
    add_bullet(doc, "Sentence-Transformers (all-MiniLM-L6-v2) para embeddings locales de 384 dimensiones sin costo y con latencia < 40ms.", "Embeddings: ")
    add_bullet(doc, "Google Gemini 1.5 Flash / OpenAI GPT-4o-mini con temperatura 0.1 y JSON Schema estructurado con Pydantic.", "Inferencia LLM: ")

    add_heading_2(doc, "2.1 Clasificador Multiclase, Resumen y Extractor de Entidades")
    ai_code = """class AIPipelineEngine:
    def classify_document(self, text_sample: str) -> dict:
        prompt = f"Clasifica en: Administrativo, Financiero o Técnico/Legal. Texto: {text_sample[:3000]}"
        res = client.models.generate_content(model="gemini-1.5-flash", contents=prompt, config={'response_mime_type': 'application/json'})
        return json.loads(res.text)

    def generate_summary(self, full_text: str) -> str:
        prompt = f"Genera resumen ejecutivo formal (150-300 palabras). Texto: {full_text[:12000]}"
        return client.models.generate_content(model="gemini-1.5-flash", contents=prompt).text.strip()

    def extract_structured_entities(self, full_text: str) -> dict:
        prompt = f"Extrae entidades (numero, emisor, fecha, monto, entidades). Texto: {full_text[:10000]}"
        res = client.models.generate_content(model="gemini-1.5-flash", contents=prompt, config={'response_mime_type': 'application/json', 'response_schema': DocumentExtractionSchema})
        return json.loads(res.text)"""
    add_code_block(doc, ai_code)

    add_heading_1(doc, "3. Motor RAG (Retrieval-Augmented Generation)")
    rag_code = """class RAGEngine:
    def answer_query(self, question: str, folder_id: str = None) -> dict:
        q_vec = embedder.encode(question).tolist()
        # Búsqueda similitud coseno en PostgreSQL con pgvector
        cursor.execute("SELECT text_content, page_number, file_name, 1 - (embedding_vector <=> %s::vector) FROM document_chunks JOIN documents ON document_chunks.document_id = documents.id ORDER BY embedding_vector <=> %s::vector LIMIT 4;", (q_vec, q_vec))
        chunks = cursor.fetchall()
        context = "\\n---\\n".join([f"[{c[2]}, Pág {c[1]}]: {c[0]}" for c in chunks if c[3] >= 0.65])
        prompt = f"Responde basándote EXCLUSIVAMENTE en el contexto:\\n{context}\\nPregunta: {question}"
        ans = client.models.generate_content(model="gemini-1.5-flash", contents=prompt).text.strip()
        return {"answer": ans, "sources": [{"file_name": c[2], "page_number": c[1]} for c in chunks if c[3] >= 0.65]}"""
    add_code_block(doc, rag_code)

    add_heading_1(doc, "4. Manejo de Casos Borde y Resiliencia")
    add_bullet(doc, "Captura de archivos corruptos sin bloquear el servicio, reintentos con Exponential Backoff ante Rate Limits (HTTP 429) y control de ventana con RecursiveSplitter.", "Resiliencia: ")

    doc.save(os.path.join(OUTPUT_DIR, "04_IA_y_Procesamiento_Documental.docx"))
    print("Dev Doc 4 generated.")

if __name__ == "__main__":
    build_doc1()
    build_doc2()
    build_doc3()
    build_doc4()
    print("All 4 Development DOCX files generated successfully!")
