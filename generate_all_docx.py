import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUTPUT_DIR = r"c:\Users\Lenovo\OneDrive\Documentos\Sistema_Inteligente_Gestion_Analisis_Documental\01_ANALISIS"

def create_base_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
    return doc

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_block(doc, title, subtitle):
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(2)
    r1 = title_p.add_run(title)
    r1.font.name = 'Arial'
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(14, 56, 122)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(12)
    r2 = sub_p.add_run(subtitle)
    r2.font.name = 'Arial'
    r2.font.size = Pt(10)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(90, 90, 90)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(14, 56, 122)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 80, 160)
    return p

def add_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.bold = True
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = RGBColor(20, 20, 20)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.bold = True
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = RGBColor(20, 20, 20)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "0E387A")
        set_cell_margins(hdr_cells[i], 100, 100, 120, 120)
        for p in hdr_cells[i].paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F4F6F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], 80, 80, 100, 100)
            for p in row_cells[c_idx].paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(40, 40, 40)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

# ==============================================================================
# DOC 1: ERS
# ==============================================================================
def build_doc1():
    doc = create_base_doc()
    add_header_block(doc, "ESPECIFICACIÓN DE REQUISITOS DE SOFTWARE (ERS)", 
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nEstándar: IEEE 830 Adaptado | Versión 1.0.0 Línea Base")

    add_heading_1(doc, "1. Control de Cambios y Metadatos de Versión")
    headers = ["Versión", "Fecha", "Autor(es)", "Rol", "Descripción del Cambio", "Estado"]
    rows = [["1.0.0", "06/09/2026", "Equipo de Ingeniería UTS", "Ingeniero de Requisitos / Arquitecto", "Creación de la Línea Base de Requisitos de Software.", "Aprobado"]]
    add_table(doc, headers, rows, [0.8, 1.0, 1.4, 1.4, 1.8, 0.8])

    add_heading_1(doc, "2. Descripción del Problema, Contexto y Justificación")
    add_heading_2(doc, "2.1 Descripción del Problema y Contexto Empresarial")
    add_p(doc, "En las Unidades Tecnológicas de Santander (UTS), los activos documentales digitales (actas, proyectos de grado, resoluciones, manuales técnicos e informes) reposan en repositorios pasivos desestructurados en formatos heterogéneos (PDF, DOCX, TXT). Esto ocasiona tiempos excesivos de búsqueda manual, falta de visibilidad ejecutiva y categorizaciones subjetivas inconsistentes.")
    
    add_heading_2(doc, "2.2 Justificación y Oportunidad de Negocio")
    add_p(doc, "La implementación de un sistema basado en Inteligencia Artificial Generativa y Semántica (RAG - Retrieval-Augmented Generation) y NLP transforma el repositorio inerte en un activo de conocimiento institucional dinámico, reduciendo en más del 70% el tiempo de consulta y consulta de información institucional clave.")

    add_heading_1(doc, "3. Objetivos del Sistema")
    add_p(doc, "Desarrollar e implementar un Sistema Inteligente de Gestión y Análisis Documental para las Unidades Tecnológicas de Santander (UTS) que centralice el almacenamiento, clasifique automáticamente archivos heterogéneos mediante IA y permita consultas semánticas en lenguaje natural con trazabilidad técnica completa.", "Objetivo General: ")
    add_bullet(doc, "Implementar un módulo seguro de autenticación JWT y control de acceso RBAC que gestione repositorios lógicos de archivos (PDF, DOCX, TXT <= 20MB).", "OE-01 (Seguridad y Almacenamiento): ")
    add_bullet(doc, "Diseñar un pipeline automatizado de ingesta, OCR y parsing para clasificar documentos en al menos 3 categorías (Administrativo, Técnico, Legal/Financiero) con precisión >= 85%.", "OE-02 (Pipeline y Clasificación): ")
    add_bullet(doc, "Incorporar servicios de IA para la generación automática de resúmenes ejecutivos y extracción de metadatos estructurados en formato JSON.", "OE-03 (Síntesis y Extracción): ")
    add_bullet(doc, "Implementar un motor RAG con base vectorial y embeddings para responder preguntas en lenguaje natural con latencia < 2.5s y citas fundamentadas.", "OE-04 (Búsqueda Semántica y RAG): ")
    add_bullet(doc, "Proveer un dashboard de analítica documental y auditoría con registro de estados (Cargado, En Proceso, Procesado, Error).", "OE-05 (Auditoría y Analítica): ")

    add_heading_1(doc, "4. Alcance del Producto y Exclusiones Explícitas")
    add_p(doc, "El producto comprende: autenticación RBAC, gestión de carpetas lógicas, carga validada de PDF/DOCX/TXT (hasta 20MB), pipeline OCR y parsing, generación de embeddings y base vectorial, clasificación multiclase, resúmenes ejecutivos, extracción estructurada JSON, consultas RAG fundamentadas, dashboard y auditoría de estados.", "Alcance Incluido: ")
    add_p(doc, "No incluye edición colaborativa de documentos en tiempo real, soporte de video/audio, firma digital externa con validez jurídica PKI, ni modificación del binario original de los documentos (inmutabilidad).", "Exclusiones Explícitas: ")

    add_heading_1(doc, "5. Requisitos Funcionales (RF-01 al RF-11)")
    rfs = [
        ("RF-01", "Autenticación y Autorización (JWT/RBAC)", "El sistema autentica credenciales y emite tokens JWT con roles (ADMIN, ANALISTA, CONSULTOR).", "Correo institucional, contraseña cifrada, rol.", "Token JWT, perfil de usuario con permisos.", "Autenticación válida retorna HTTP 200 en <500ms; error retorna HTTP 401."),
        ("RF-02", "Gestión de Repositorios y Carpetas", "Permite crear, listar, renombrar y estructurar carpetas jerárquicas lógicas.", "Nombre de carpeta, carpeta padre ID, permisos.", "UUID de carpeta, ruta lógica normalizada.", "Creación sin nombres duplicados en mismo nivel en <300ms."),
        ("RF-03", "Carga y Validación de Archivos", "Recibe archivos PDF, DOCX y TXT validando Magic Bytes y límite de 20 MB.", "Binario multipart, carpeta destino, usuario.", "Registro con estado 'Cargado', SHA-256 e ID.", "Rechazo inmediato si excede 20MB o formato inválido (HTTP 400/413)."),
        ("RF-04", "Extracción de Texto y Parsing / OCR", "Extrae texto nativo o ejecuta OCR en PDFs escaneados y genera chunks solapados.", "Archivo binario almacenado, config chunking.", "Texto plano normalizado, chunks con metadatos.", "Extracción completa preservando orden; estado 'En Proceso'."),
        ("RF-05", "Clasificación Automática Multiclase", "Clasifica automáticamente el archivo en Administrativo, Técnico o Legal/Financiero.", "Chunks y resumen inicial del documento.", "Categoría asignada, score confianza (0-1), justificación.", "Precisión del modelo >= 85% sobre conjunto de validación."),
        ("RF-06", "Generación de Resumen Ejecutivo", "Genera síntesis ejecutiva concisa (150-300 palabras) de puntos clave y conclusiones.", "Texto representativo del documento.", "Resumen en Markdown estructurado.", "Texto coherente sin alucinaciones persistido en base de datos."),
        ("RF-07", "Extracción de Entidades y JSON Clave", "Extrae entidades estructuradas (responsables, fechas, montos, resoluciones) en JSON.", "Texto del documento, esquema JSON.", "JSON estructurado validado contra esquema.", "JSON válido contra schema en 95% de casos evaluados."),
        ("RF-08", "Consulta Semántica RAG en Lenguaje Natural", "Responde preguntas en lenguaje natural recuperando fragmentos vectoriales con citas.", "Pregunta del usuario, filtros opcionales.", "Respuesta redactada con citas [Documento, Página].", "Respuesta fundamentada en contexto documental en < 2.5s."),
        ("RF-09", "Búsqueda Tradicional y Filtros", "Búsqueda léxica por palabras clave, fechas, categorías y metadatos.", "Término de búsqueda, filtros booleanos.", "Lista paginada de documentos con resaltado.", "Retorno en <300ms para bases de hasta 50.000 registros."),
        ("RF-10", "Dashboard de Analítica y Métricas", "Gráficos de volumen, distribución por categorías, tasas de éxito y tiempos medios.", "Rango de fechas y filtros temporales.", "Tarjetas de KPI, gráficos interactivos.", "Carga en <1.5s con actualización reactiva."),
        ("RF-11", "Registro y Auditoría de Procesamiento", "Registra transiciones de estado (Cargado -> En Proceso -> Procesado/Error) con trazas.", "Eventos de procesamiento y excepciones.", "Log transaccional inmutable con timestamps.", "Trazabilidad del 100% de documentos con histórico.")
    ]
    for code, name, desc, ent, sal, crit in rfs:
        add_heading_2(doc, f"{code}: {name}")
        add_p(doc, desc, "Descripción: ")
        add_p(doc, ent, "Entradas: ")
        add_p(doc, sal, "Salidas: ")
        add_p(doc, crit, "Criterio de Aceptación: ")

    add_heading_1(doc, "6. Requisitos No Funcionales (RNF-01 al RNF-06)")
    rnf_headers = ["Código", "Categoría", "Especificación Técnica Detallada", "Criterio de Aceptación / SLA"]
    rnf_rows = [
        ["RNF-01", "Seguridad", "Cifrado BCrypt (cost>=12), TLS 1.3, JWT RSA256, Magic Bytes validation y sanitización de inputs.", "Cero vulnerabilidades críticas SAST/DAST."],
        ["RNF-02", "Rendimiento", "Latencia P95 <= 2.5s en consultas RAG; <= 300ms en endpoints transaccionales CRUD.", "50 usuarios concurrentes sin degradación."],
        ["RNF-03", "Disponibilidad", "Disponibilidad del 99.5% en horario de operación académica institucional.", "Inactividad no programada < 3.6 horas/mes."],
        ["RNF-04", "Usabilidad", "Interfaz web responsiva con directrices WCAG 2.1 AA y heurísticas de Nielsen.", "Puntuación System Usability Scale >= 80/100."],
        ["RNF-05", "Mantenibilidad", "Clean Architecture, principios SOLID, documentación OpenAPI 3.0, cobertura pruebas >=80%.", "SonarQube grado 'A', JaCoCo >= 80%."],
        ["RNF-06", "Escalabilidad", "Arquitectura desacoplada en contenedores Docker lista para escalamiento horizontal de workers.", "Hasta 500 docs/hora escalando réplicas IA."]
    ]
    add_table(doc, rnf_headers, rnf_rows, [0.8, 1.1, 3.2, 1.9])

    add_heading_1(doc, "7. Reglas de Negocio (RN-01 a RN-05)")
    add_bullet(doc, "Archivos con peso > 20MB o Magic Bytes incompatibles con PDF/DOCX/TXT son rechazados en el API Gateway.", "RN-01 (Restricción de Formato): ")
    add_bullet(doc, "El procesamiento de IA y OCR debe ser asíncrono mediante colas de tareas sin bloquear hilos HTTP.", "RN-02 (Concurrencia y Encolamiento): ")
    add_bullet(doc, "Reintentos exponenciales automáticos (hasta 3 intentos) ante respuestas HTTP 429/503 del proveedor de IA.", "RN-03 (Política de Reintentos): ")
    add_bullet(doc, "Aislamiento estricto de documentos en base vectorial según permisos RBAC del usuario consultante.", "RN-04 (Gobernanza y Privacidad): ")
    add_bullet(doc, "Los documentos en estado 'Procesado' son inmutables; cambios requieren creación de nueva versión vinculada.", "RN-05 (Inmutabilidad Documental): ")

    doc.save(os.path.join(OUTPUT_DIR, "01_ERS_Especificacion_Requisitos_Software.docx"))
    print("Doc 1 generated.")

# ==============================================================================
# DOC 2: HISTORIAS DE USUARIO
# ==============================================================================
def build_doc2():
    doc = create_base_doc()
    add_header_block(doc, "DOCUMENTO DE HISTORIAS DE USUARIO (HU)",
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nMetodología Ágil (Scrum / INVEST / BDD Gherkin)")

    add_heading_1(doc, "1. Matriz Resumen de Historias de Usuario")
    headers = ["ID", "Épica", "Título de la Historia", "Rol Principal", "Story Points", "Prioridad"]
    rows = [
        ["HU-01", "Seguridad y Acceso", "Inicio de Sesión Seguro y Gestión de Tokens", "Usuario General", "3", "Alta (Must)"],
        ["HU-02", "Gestión Documental", "Organización y Estructuración en Carpetas Lógicas", "Analista Documental", "5", "Alta (Must)"],
        ["HU-03", "Ingesta de Archivos", "Carga y Validación de Archivos Heterogéneos", "Analista Documental", "5", "Alta (Must)"],
        ["HU-04", "Inteligencia Artificial", "Clasificación Automática Multiclase de Documentos", "Analista Documental", "8", "Alta (Must)"],
        ["HU-05", "Inteligencia Artificial", "Generación Automatizada de Resúmenes Ejecutivos", "Consultor / Analista", "5", "Media (Should)"],
        ["HU-06", "Inteligencia Artificial", "Extracción Estructurada de Entidades y Metadatos", "Analista / Auditor", "8", "Alta (Must)"],
        ["HU-07", "Búsqueda y RAG", "Consulta Semántica en Lenguaje Natural (Q&A)", "Consultor / Auditor", "13", "Alta (Must)"],
        ["HU-08", "Analítica y Auditoría", "Visualización de Dashboard de Métricas y Estado", "Administrador", "5", "Media (Should)"]
    ]
    add_table(doc, headers, rows, [0.7, 1.5, 2.3, 1.3, 0.6, 0.8])

    add_heading_1(doc, "2. Especificación Detallada de Historias de Usuario (BDD / Gherkin)")

    hus = [
        ("HU-01", "Inicio de Sesión Seguro y Gestión de Tokens", "Seguridad y Acceso", "3", "Alta (Must)",
         "Como usuario del sistema (Administrador, Analista o Consultor), quiero autenticarme con mis credenciales institucionales seguras, para acceder a las funcionalidades del sistema según los permisos asignados a mi rol.",
         "Dado que el usuario se encuentra en la pantalla de login con credenciales activas ('analista.uts@uts.edu.co'),\nCuando ingresa correo y contraseña válidos y presiona 'Iniciar Sesión',\nEntonces el sistema valida la identidad en <500ms, emite token JWT con rol 'ANALISTA' y redirige a la vista principal.",
         "Dado que el usuario ingresa una contraseña incorrecta o correo inexistente,\nCuando envía el formulario de inicio de sesión,\nEntonces el sistema no emite token, responde HTTP 401 y muestra 'Credenciales incorrectas' sin dar detalles adicionales."),

        ("HU-02", "Organización y Estructuración en Carpetas Lógicas", "Gestión Documental", "5", "Alta (Must)",
         "Como Analista de Información Documental, quiero crear y estructurar carpetas jerárquicas en el repositorio, para categorizar y segregar los archivos institucionales de manera ordenada.",
         "Dado que el Analista tiene permisos de escritura en el módulo documental,\nCuando crea una carpeta con el nombre 'Proyectos_Grado_2026',\nEntonces el sistema genera su UUID, la agrega al árbol de directorios y persiste el registro en la base de datos.",
         "Dado que ya existe una carpeta llamada 'Proyectos_Grado_2026' en el mismo nivel jerárquico,\nCuando el usuario intenta crear otra con idéntico nombre,\nEntonces el sistema rechaza la solicitud y muestra 'Ya existe una carpeta con este nombre en la ubicación actual'."),

        ("HU-03", "Carga y Validación de Archivos Heterogéneos", "Ingesta de Archivos", "5", "Alta (Must)",
         "Como Analista de Información Documental, quiero cargar archivos individuales o por lotes en formatos PDF, DOCX y TXT, para que queden almacenados de forma segura e inicien el pipeline de análisis inteligente.",
         "Dado que el usuario selecciona 'resolucion_rectoral.pdf' de 4.5 MB hacia la carpeta 'Actas',\nCuando presiona 'Cargar Archivo',\nEntonces el sistema valida Magic Bytes (%PDF), peso <= 20MB, persiste el archivo con hash SHA-256, marca estado 'Cargado' y encola el procesamiento asíncrono.",
         "Dado que el usuario intenta subir un archivo de 25.8 MB,\nCuando presiona 'Cargar Archivo',\nEntonces el sistema intercepta la petición, devuelve HTTP 413 y notifica 'El archivo excede el tamaño máximo permitido de 20 MB'."),

        ("HU-04", "Clasificación Automática Multiclase de Documentos", "Inteligencia Artificial", "8", "Alta (Must)",
         "Como Analista Documental, quiero que el sistema clasifique automáticamente los documentos cargados en categorías como Administrativo, Técnico o Legal/Financiero, para eliminar la clasificación manual y asegurar la homogeneidad del archivo.",
         "Dado que un documento con contenido sobre presupuesto y comités fue parseado exitosamente,\nCuando el clasificador de IA procesa los fragmentos,\nEntonces asigna la categoría 'Administrativo' con confianza >= 85% y actualiza los metadatos del documento.",
         "Dado que un documento presenta texto disperso con confianza menor a 0.60,\nCuando concluye la inferencia,\nEntonces se etiqueta como 'Por Clasificar / Revisión Requerida' y se genera alerta visual en la bandeja de revisión."),

        ("HU-05", "Generación Automatizada de Resúmenes Ejecutivos", "Inteligencia Artificial", "5", "Media (Should)",
         "Como Consultor o Directivo Institucional, quiero disponer de un resumen ejecutivo sintetizado de cada documento cargado, para comprender el propósito y puntos clave del archivo sin necesidad de leerlo por completo.",
         "Dado que un informe técnico de 35 páginas fue parseado correctamente,\nCuando el worker ejecuta la rutina de síntesis con el LLM,\nEntonces genera un resumen en Markdown de 150-300 palabras con propósito, puntos clave y conclusiones.",
         "Dado que se procesa un archivo con menos de 20 palabras legibles,\nCuando se intenta generar el resumen,\nEntonces el sistema detecta longitud insuficiente y marca 'Resumen no disponible - Contenido insuficiente' sin gastar llamadas al LLM."),

        ("HU-06", "Extracción Estructurada de Entidades y Metadatos", "Inteligencia Artificial", "8", "Alta (Must)",
         "Como Auditor o Analista de la UTS, quiero que el sistema extraiga entidades clave (fechas, responsables, valores económicos, resoluciones) en formato JSON estructurado, para facilitar auditorías y exportación de datos.",
         "Dado que el documento corresponde a un contrato institucional legal/financiero,\nCuando el modelo de IA procesa la extracción según el esquema JSON,\nEntonces devuelve objeto JSON validado con: numero_contrato, fecha_firma, partes, cuantia y vigencia.",
         "Dado que el contrato no especifica cuantía económica,\nCuando el extractor genera el JSON,\nEntonces asigna 'null' a dicho atributo y pasa la validación del esquema Pydantic sin generar excepciones."),

        ("HU-07", "Consulta Semántica en Lenguaje Natural (Q&A con RAG)", "Búsqueda y RAG", "13", "Alta (Must)",
         "Como Consultor o Investigador institucional, quiero hacer preguntas en lenguaje natural sobre la base documental institucional, para obtener respuestas precisas fundamentadas y sustentadas con citas exactas de los documentos fuente.",
         "Dado que existen documentos indexados en la base de datos vectorial,\nCuando el usuario pregunta: '¿Cuáles son los requisitos de graduación según el reglamento vigente?',\nEntonces el motor recupera los chunks más similares, el LLM sintetiza la respuesta y entrega la cita [Reglamento_Estudiantil_UTS.pdf, Pág. 18] en <2.5s.",
         "Dado que el usuario pregunta sobre un tema no contenido en el repositorio,\nCuando el motor RAG calcula similitudes y quedan bajo el umbral (score < 0.65),\nEntonces el sistema responde: 'La información solicitada no se encuentra disponible en los documentos cargados' sin alucinar."),

        ("HU-08", "Visualización de Dashboard de Métricas y Estado", "Analítica y Auditoría", "5", "Media (Should)",
         "Como Administrador del Sistema, quiero consultar un dashboard interactivo con estadísticas de carga, distribución de categorías y estado de auditoría, para monitorear la salud del sistema y la actividad documental de las UTS.",
         "Dado que el Administrador accede a la sección de Analítica,\nCuando la vista carga en <1.5s,\nEntonces muestra gráficos por categoría (Administrativo, Técnico, Legal), total de documentos, tasa de éxito y latencias medias.",
         "Dado que el Administrador filtra por el rango de fechas del último mes,\nCuando aplica el filtro,\nEntonces las gráficas y contadores se actualizan dinámicamente vía API sin recarga de página.")
    ]

    for hid, title, epica, sp, prio, decl, esc1, esc2 in hus:
        add_heading_2(doc, f"{hid}: {title} (Épica: {epica} | {sp} SP | Prioridad: {prio})")
        add_p(doc, decl, "Declaración: ")
        add_p(doc, esc1, "Escenario 1 (Éxito): ")
        add_p(doc, esc2, "Escenario 2 (Excepción / Límite): ")

    doc.save(os.path.join(OUTPUT_DIR, "02_Historias_de_Usuario.docx"))
    print("Doc 2 generated.")

# ==============================================================================
# DOC 3: CASOS DE USO
# ==============================================================================
def build_doc3():
    doc = create_base_doc()
    add_header_block(doc, "ESPECIFICACIÓN DE CASOS DE USO DEL SISTEMA",
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nEstándar UML / Plantilla RUP - Alistair Cockburn")

    add_heading_1(doc, "1. Tabla Resumen de Casos de Uso")
    headers = ["Código", "Nombre del Caso de Uso", "Actor Principal", "Actores Secundarios", "Complejidad"]
    rows = [
        ["CU-01", "Autenticar Usuario e Inicializar Sesión", "Usuario General", "Servidor Auth / JWT", "Baja"],
        ["CU-02", "Cargar y Parsear Documento al Repositorio", "Analista Documental", "Motor OCR / Storage", "Media"],
        ["CU-03", "Ejecutar Pipeline de Análisis e Inferencia de IA", "Sistema (Worker IA)", "Proveedor LLM / DB", "Alta"],
        ["CU-04", "Realizar Consulta Semántica Documental (RAG)", "Consultor / Auditor", "Motor Vectorial / LLM", "Alta"],
        ["CU-05", "Gestionar Carpetas y Estructura Jerárquica", "Analista Documental", "Base de Datos", "Baja"],
        ["CU-06", "Realizar Búsqueda Léxica y Filtrado", "Consultor / Analista", "Base de Datos", "Baja"],
        ["CU-07", "Consultar Dashboard y Métricas de Rendimiento", "Administrador", "Servicio de Analítica", "Media"],
        ["CU-08", "Auditar Logs y Trazas de Error del Pipeline", "Administrador", "Log Storage / DB", "Media"]
    ]
    add_table(doc, headers, rows, [0.8, 2.5, 1.4, 1.4, 0.9])

    add_heading_1(doc, "2. Especificación Formal de los 4 Casos de Uso Centrales")

    cus = [
        ("CU-01: Autenticar Usuario e Inicializar Sesión",
         "Usuario (ADMIN, ANALISTA, CONSULTOR)", "Servicio JWT / Spring Security",
         "Usuario activo y registrado en la base de datos institucional.",
         "Usuario ingresa a la aplicación y envía el formulario de inicio de sesión.",
         ["1. El usuario navega a la URL de login e ingresa correo institucional y contraseña.",
          "2. El Frontend valida formato y despacha POST /api/v1/auth/login.",
          "3. El Backend busca el usuario, valida hash BCrypt y verifica estado ACTIVO.",
          "4. El Backend genera Access Token y Refresh Token JWT con los claims del rol.",
          "5. El Backend retorna HTTP 200 y el Frontend redirige al panel según el rol."],
         ["FA-01: Renovación con Refresh Token cuando el Access Token expira.",
          "FA-02: Redirección diferenciada (ADMIN a Auditoría, ANALISTA a Repositorio, CONSULTOR a Chat RAG)."],
         ["FE-01: Credenciales inválidas -> HTTP 401 Unauthorized.",
          "FE-02: Usuario inactivo -> HTTP 403 Forbidden.",
          "FE-03: Fallo de base de datos -> HTTP 503 Service Unavailable."],
         "Éxito: Sesión activa y token JWT emitido. Fallo: Sin sesión y evento registrado en auditoría."),

        ("CU-02: Cargar y Parsear Documento al Repositorio",
         "Analista de Información Documental", "Storage Local / DB / Cola de Tareas",
         "Usuario autenticado con rol ADMIN o ANALISTA y carpeta destino válida.",
         "Usuario arrastra o selecciona archivo (PDF, DOCX, TXT) y confirma subida.",
         ["1. El usuario selecciona la carpeta destino y presiona 'Subir Archivo'.",
          "2. El Frontend valida extensión y peso <= 20MB y despacha POST /api/v1/documents/upload.",
          "3. El Backend valida Magic Bytes (%PDF, PK.., etc.) y calcula hash SHA-256.",
          "4. El Backend persiste el binario en almacenamiento seguro y crea registro con estado 'Cargado'.",
          "5. El Backend encola el evento de procesamiento asíncrono y responde HTTP 201 Created."],
         ["FA-01: Carga masiva por lote procesando archivos de forma iterativa.",
          "FA-02: Carga de TXT plano pasando directamente a fase de chunking."],
         ["FE-01: Extensión o Magic Bytes no permitidos -> HTTP 400 Bad Request.",
          "FE-02: Archivo > 20 MB -> HTTP 413 Payload Too Large.",
          "FE-03: Fallo de disco/storage -> Rollback transaccional y HTTP 500."],
         "Éxito: Documento guardado inmutable con estado 'Cargado' en cola. Fallo: Archivo descartado."),

        ("CU-03: Ejecutar Pipeline de Análisis e Inferencia de IA",
         "Sistema (Worker Asíncrono FastAPI/Celery)", "OCR Tesseract, Embeddings, Vector DB, LLM",
         "Documento en estado 'Cargado' y binario accesible en storage.",
         "Llegada de mensaje a la cola de procesamiento de documentos.",
         ["1. El Worker toma la tarea y actualiza el estado a 'En Proceso'.",
          "2. Invoca el módulo de extracción (PyMuPDF / python-docx / OCR Tesseract si no hay texto digital).",
          "3. Limpia y divide el texto en chunks (500 tokens con overlap de 100).",
          "4. Genera embeddings vectoriales e indexa en Vector DB (Qdrant/pgvector).",
          "5. Ejecuta modelo de clasificación multiclase (Administrativo, Técnico, Legal/Financiero).",
          "6. Invoca prompt de síntesis en LLM para generar Resumen Ejecutivo.",
          "7. Invoca prompt estructurado con JSON Schema para extraer entidades clave.",
          "8. Actualiza base de datos con estado 'Procesado', resumen y JSON de entidades."],
         ["FA-01: Activación automática de OCR cuando el PDF contiene sólo imágenes escaneadas.",
          "FA-02: Reintentos con Exponential Backoff ante errores de rate limit (HTTP 429)."],
         ["FE-01: Archivo corrupto o ilegible -> Estado 'Error' y log con traza detallada.",
          "FE-02: Caída persistente del proveedor de IA tras 3 reintentos -> Estado 'Error'."],
         "Éxito: Estado 'Procesado', vectores indexados y metadatos persistidos. Fallo: Estado 'Error'."),

        ("CU-04: Realizar Consulta Semántica Documental (RAG)",
         "Consultor / Auditor / Analista", "Motor Embeddings, Vector DB, LLM Generativo",
         "Usuario autenticado y documentos en estado 'Procesado' accesibles según rol.",
         "Usuario escribe pregunta en lenguaje natural y presiona 'Consultar'.",
         ["1. El usuario formula su consulta (ej. '¿Cuáles son los plazos de entrega del informe?').",
          "2. El Frontend despacha POST /api/v1/rag/query con la pregunta y filtros.",
          "3. El módulo de IA genera el embedding de la pregunta.",
          "4. Busca por similitud de coseno en Vector DB con filtros RBAC recuperando Top-K chunks (K=4).",
          "5. Inyecta los fragmentos en el System Prompt instruyendo al LLM a responder basado en el contexto.",
          "6. El LLM genera la respuesta con citas explícitas [Documento, Página].",
          "7. El Backend retorna la respuesta en <2.5s y la interfaz muestra tarjetas interactivas."],
         ["FA-01: Búsqueda restringida a una única carpeta o documento seleccionado.",
          "FA-02: Conversación contextual multiturno incorporando historial previo."],
         ["FE-01: Similitud inferior a umbral (score < 0.65) -> Retorna 'Información no encontrada'.",
          "FE-02: Timeout de Vector DB (> 2.0s) -> HTTP 504 con mensaje controlado."],
         "Éxito: Respuesta fundamentada con citas entregada en <2.5s. Fallo: Mensaje de contingencia.")
    ]

    for title, act_p, act_s, pre, trig, steps, fas, fes, post in cus:
        add_heading_2(doc, title)
        add_p(doc, act_p, "Actor Principal: ")
        add_p(doc, act_s, "Actores Secundarios: ")
        add_p(doc, pre, "Precondiciones: ")
        add_p(doc, trig, "Disparador (Trigger): ")
        add_p(doc, "", "Flujo Principal:")
        for s in steps:
            add_bullet(doc, s)
        add_p(doc, "", "Flujos Alternativos:")
        for fa in fas:
            add_bullet(doc, fa)
        add_p(doc, "", "Flujos de Excepción:")
        for fe in fes:
            add_bullet(doc, fe)
        add_p(doc, post, "Postcondiciones: ")

    doc.save(os.path.join(OUTPUT_DIR, "03_Casos_de_Uso_y_Especificaciones.docx"))
    print("Doc 3 generated.")

# ==============================================================================
# DOC 4: PERFILES DE USUARIO Y PERSONAS
# ==============================================================================
def build_doc4():
    doc = create_base_doc()
    add_header_block(doc, "PERFILES DE USUARIO, PERSONAS Y ANÁLISIS DE STAKEHOLDERS",
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nEstándar: User Personas & Matriz RACI")

    add_heading_1(doc, "1. Mapa de Actores del Sistema")
    add_bullet(doc, "Responsable de la infraestructura, seguridad, observabilidad, gestión de cuentas y control de costos de APIs de IA.", "Administrador del Sistema: ")
    add_bullet(doc, "Responsable operativo de la ingesta documental masiva, organización de carpetas y verificación de la clasificación taxonómica y entidades JSON.", "Analista de Información Documental: ")
    add_bullet(doc, "Usuario final que realiza consultas semánticas en lenguaje natural, consulta resúmenes ejecutivos y audita evidencias documentales.", "Consultor / Auditor Institucional: ")

    add_heading_1(doc, "2. Fichas Técnicas de User Persona")

    personas = [
        ("Ing. Carlos Mendoza Villamizar (38 años)", "Administrador de Infraestructura TI y Seguridad - UTS", "Avanzado (Senior - Linux, Docker, PostgreSQL, APIs, Seguridad)",
         "Garantizar disponibilidad del 99.5%, monitorear costos y tokens de IA, auditar logs de error y gestionar accesos RBAC.",
         "Logs crípticos en fallos asíncronos, reportes tardíos de caídas de servicios y falta de observabilidad.",
         "CU-01 (Login), CU-07 (Dashboard Métricas), CU-08 (Auditoría de Logs).",
         "'Necesito que la plataforma sea blindada y que cada falla en el pipeline de IA quede registrada con su causa raíz.'"),

        ("Laura Patricia Gómez Serrano (29 años)", "Asistente de Gestión Documental y Archivo Central - UTS", "Intermedio (Manejo de repositorios cloud, ofimática avanzada y gestores documentales)",
         "Cargar cientos de documentos rápidamente, asegurar que la clasificación automática sea correcta y extraer metadatos estructurados en JSON.",
         "Extracciones fallidas sin aviso, interfaces lentas y tener que clasificar manualmente actas o contratos extensos.",
         "CU-02 (Carga de Archivos), CU-03 (Monitoreo de Pipeline), CU-05 (Gestión de Carpetas).",
         "'Subo cientos de resoluciones al mes. La extracción de fechas y números de acta me ahorra semanas de trabajo.'"),

        ("Dr. Fernando Ruiz Beltrán (46 años)", "Par Académico / Auditor Interno de Calidad - UTS", "Básico - Medio (Usuario de navegadores web y aplicaciones de consulta)",
         "Encontrar respuestas normativas de inmediato, leer resúmenes ejecutivos y verificar citas documentales exactas.",
         "Búsquedas tradicionales que devuelven cientos de PDFs sin indicar en qué página está la respuesta, y alucinaciones de IA.",
         "CU-04 (Consulta Semántica RAG), CU-06 (Búsqueda Léxica), Lectura de Resúmenes Ejecutivos.",
         "'Quiero preguntarle al sistema en lenguaje natural y que me muestre el párrafo exacto que sustenta la respuesta.'")
    ]

    for name, rol, skill, obj, frust, cus_f, cita in personas:
        add_heading_2(doc, f"User Persona: {name}")
        add_p(doc, rol, "Rol Institucional: ")
        add_p(doc, skill, "Nivel Tecnológico: ")
        add_p(doc, obj, "Objetivos en el Software: ")
        add_p(doc, frust, "Frustraciones Habituales: ")
        add_p(doc, cus_f, "Casos de Uso Frecuentes: ")
        add_p(doc, cita, "Cita Representativa: ")

    add_heading_1(doc, "3. Matriz RACI de Interacción Funcional")
    raci_headers = ["Módulo / Funcionalidad", "Administrador TI", "Analista Documental", "Consultor / Auditor", "Worker IA / Backend"]
    raci_rows = [
        ["M01: Autenticación y RBAC", "A / R", "I", "I", "R"],
        ["M02: Gestión de Carpetas", "A", "R", "I", "I"],
        ["M03: Carga y Validación de Archivos", "I", "A / R", "I", "R"],
        ["M04: Extracción de Texto y OCR", "I", "C", "I", "A / R"],
        ["M05: Clasificación Multiclase IA", "I", "A", "I", "R"],
        ["M06: Resúmenes y JSON Estructurado", "I", "A", "C", "R"],
        ["M07: Búsqueda Semántica y RAG", "I", "C", "A / R", "R"],
        ["M08: Dashboard y Métricas", "A / R", "I", "I", "R"],
        ["M09: Auditoría y Manejo de Errores", "A / R", "I", "I", "R"]
    ]
    add_table(doc, raci_headers, raci_rows, [2.5, 1.1, 1.2, 1.2, 1.2])

    doc.save(os.path.join(OUTPUT_DIR, "04_Perfiles_de_Usuario_y_Personas.docx"))
    print("Doc 4 generated.")

# ==============================================================================
# DOC 5: PRIORIZACIÓN DE REQUISITOS
# ==============================================================================
def build_doc5():
    doc = create_base_doc()
    add_header_block(doc, "PRIORIZACIÓN DE REQUISITOS DEL SOFTWARE",
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nTécnicas: MoSCoW & Matriz de Valor vs. Complejidad")

    add_heading_1(doc, "1. Justificación de la Técnica Seleccionada")
    add_p(doc, "La combinación de MoSCoW y la Matriz de Valor vs. Complejidad permite blindar el Producto Mínimo Viable (MVP) asegurando que los requerimientos de mayor retorno institucional (RAG, clasificación multiclase, ingesta segura) se desarrollen con máxima prioridad, optimizando el esfuerzo de desarrollo y gestionando la incertidumbre técnica de los componentes de Inteligencia Artificial.")

    add_heading_1(doc, "2. Clasificación MoSCoW de Requisitos")
    headers = ["Categoría MoSCoW", "Código RF", "Nombre del Requisito", "Justificación Estratégica"]
    rows = [
        ["MUST HAVE", "RF-01", "Autenticación y RBAC", "Seguridad y control de acceso indispensable para entorno institucional."],
        ["MUST HAVE", "RF-02", "Gestión de Carpetas", "Estructuración indispensable del repositorio documental."],
        ["MUST HAVE", "RF-03", "Carga y Validación de Archivos", "Ingesta base para alimentar el pipeline analítico."],
        ["MUST HAVE", "RF-04", "Extracción de Texto y OCR", "Precondición técnica ineludible para la inferencia de IA."],
        ["MUST HAVE", "RF-05", "Clasificación Multiclase", "Núcleo de catalogación automática del repositorio."],
        ["MUST HAVE", "RF-07", "Extracción de Entidades JSON", "Estructuración de datos para cruces y auditoría."],
        ["MUST HAVE", "RF-08", "Consulta Semántica RAG", "Diferenciador principal y funcionalidad core del proyecto."],
        ["SHOULD HAVE", "RF-06", "Resumen Ejecutivo", "Alto valor gerencial para toma rápida de decisiones."],
        ["SHOULD HAVE", "RF-09", "Búsqueda Tradicional", "Alternativa de recuperación complementaria por palabras clave."],
        ["SHOULD HAVE", "RF-11", "Auditoría de Procesamiento", "Trazabilidad técnica para resolución de incidencias."],
        ["COULD HAVE", "RF-10", "Dashboard de Analítica", "Visualización de métricas complementaria para monitoreo."],
        ["WON'T HAVE", "RF-EX01", "Edición Colaborativa en Vivo", "Fuera del alcance del sistema gestor y analizador."],
        ["WON'T HAVE", "RF-EX02", "Firma Digital Criptográfica", "Requiere integración externa con entidad certificadora PKI."]
    ]
    add_table(doc, headers, rows, [1.3, 0.8, 2.2, 2.9])

    add_heading_1(doc, "3. Matriz de Ponderación (Valor vs. Complejidad)")
    mat_headers = ["Código RF", "Requisito Funcional", "Valor Negocio (1-5)", "Complejidad (1-5)", "Ratio (V/C)", "Cuadrante"]
    mat_rows = [
        ["RF-01", "Autenticación JWT / RBAC", "5", "2", "2.50", "Victoria Rápida (Quick Win)"],
        ["RF-02", "Gestión de Carpetas", "4", "2", "2.00", "Victoria Rápida (Quick Win)"],
        ["RF-03", "Carga y Validación Archivos", "5", "2", "2.50", "Victoria Rápida (Quick Win)"],
        ["RF-04", "Extracción Texto y OCR", "5", "4", "1.25", "Proyecto Estratégico (Core)"],
        ["RF-05", "Clasificación Multiclase", "5", "4", "1.25", "Proyecto Estratégico (Core)"],
        ["RF-06", "Resumen Ejecutivo", "4", "3", "1.33", "Alto Valor / Media Complejidad"],
        ["RF-07", "Extracción Entidades JSON", "5", "4", "1.25", "Proyecto Estratégico (Core)"],
        ["RF-08", "Consulta Semántica RAG", "5", "5", "1.00", "Estratégico Mayor (High Impact)"],
        ["RF-09", "Búsqueda Tradicional", "3", "2", "1.50", "Victoria Rápida (Quick Win)"],
        ["RF-10", "Dashboard Analítica", "3", "3", "1.00", "Complementario"],
        ["RF-11", "Auditoría de Procesamiento", "4", "3", "1.33", "Infraestructura / Calidad"]
    ]
    add_table(doc, mat_headers, mat_rows, [0.8, 2.2, 1.1, 1.1, 0.8, 1.2])

    add_heading_1(doc, "4. Cronograma Preliminar de Entregas por Sprint")
    add_bullet(doc, "RF-01, RF-02, RF-03, RF-11 (Base). Entregables: Seguridad JWT/RBAC, API y UI de carpetas y subida de archivos, base relacional PostgreSQL.", "Sprint 1 (Semanas 1-2) - Cimientos y Gestión de Archivos: ")
    add_bullet(doc, "RF-04, RF-05, RF-11 (Logs avanzados). Entregables: Microservicio IA FastAPI/Celery, parsing multiformato, OCR Tesseract y clasificador multiclase.", "Sprint 2 (Semanas 3-4) - Pipeline de Extracción y Clasificación: ")
    add_bullet(doc, "RF-06, RF-07, RF-08. Entregables: Resúmenes ejecutivos, extracción JSON, indexación en base vectorial y módulo Q&A conversacional RAG.", "Sprint 3 (Semanas 5-6) - Extracción de Entidades y Motor RAG: ")
    add_bullet(doc, "RF-09, RF-10, Estabilización RNF. Entregables: Dashboard analítico, búsqueda combinada, pruebas de carga y optimización de latencias (<2.5s).", "Sprint 4 (Semanas 7-8) - Dashboard, Búsqueda y Estabilización: ")

    doc.save(os.path.join(OUTPUT_DIR, "05_Priorizacion_de_Requisitos.docx"))
    print("Doc 5 generated.")

# ==============================================================================
# DOC 6: ANÁLISIS DE RIESGOS
# ==============================================================================
def build_doc6():
    doc = create_base_doc()
    add_header_block(doc, "ANÁLISIS Y GESTIÓN DE RIESGOS TÉCNICOS Y DE IA",
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nEstándar: ISO 31000 / OWASP Top 10 for LLM Applications")

    add_heading_1(doc, "1. Metodología de Evaluación de Riesgos")
    add_p(doc, "La evaluación cuantitativa se rige por la fórmula: Severidad = Probabilidad (P: 1 a 5) x Impacto (I: 1 a 5). Los niveles se clasifican en: Bajo (1 a 6 - Monitoreo), Medio (8 a 12 - Mitigación programada) y Alto (15 a 25 - Atención inmediata y plan de contingencia obligatorio).")

    add_heading_1(doc, "2. Tabla Maestra de Riesgos Técnicos y de IA")
    headers = ["ID", "Riesgo Identificado", "Categoría", "P (1-5)", "I (1-5)", "Severidad (PxI)", "Nivel"]
    rows = [
        ["RSK-01", "Alucinaciones o imprecisiones en respuestas IA", "Inteligencia Artificial", "4", "4", "16", "Alto"],
        ["RSK-02", "Latencias elevadas por embeddings u OCR", "Rendimiento", "4", "3", "12", "Medio"],
        ["RSK-03", "Archivos corruptos o PDFs escaneados borrosos", "Calidad de Datos", "3", "4", "12", "Medio"],
        ["RSK-04", "Agotamiento de cuotas / sobrecostos de IA", "Costos / Infraestructura", "3", "4", "12", "Medio"],
        ["RSK-05", "Inyección indirecta de prompts en documentos", "Ciberseguridad", "3", "5", "15", "Alto"],
        ["RSK-06", "Fuga de información entre roles en búsqueda RAG", "Seguridad / Privacidad", "2", "5", "10", "Medio"]
    ]
    add_table(doc, headers, rows, [0.7, 2.5, 1.5, 0.5, 0.5, 0.8, 0.7])

    add_heading_1(doc, "3. Planes de Acción, Mitigación y Contingencia")

    rsks = [
        ("RSK-01: Alucinaciones o imprecisiones en respuestas IA",
         "Mitigación: Arquitectura RAG estricta con prompt negativo rígido ('Responde únicamente con el contexto provisto...'), temperatura LLM <= 0.1 y citas obligatorias de origen.",
         "Contingencia: Si el score vectorial de similitud es < 0.65, el sistema no invoca al LLM y responde: 'Información no disponible en el repositorio institucional'."),

        ("RSK-02: Latencias elevadas por embeddings u OCR",
         "Mitigación: Procesamiento asíncrono con colas Celery/Redis, caché de embeddings en Redis para consultas frecuentes y segmentación en memoria.",
         "Contingencia: Timeouts estrictos (2.0s vector search, 5.0s LLM) y degradación elegante a búsqueda léxica tradicional si el motor RAG no responde a tiempo."),

        ("RSK-03: Archivos corruptos o PDFs escaneados borrosos",
         "Mitigación: Prevalidación de cabeceras con PyMuPDF/python-docx y fallback automático a OCR Tesseract con binarización y realce de contraste OpenCV.",
         "Contingencia: Captura controlada de excepciones, marcado en estado 'Error' sin detener el worker, registro de traza y notificación al usuario."),

        ("RSK-04: Agotamiento de cuotas / sobrecostos de IA",
         "Mitigación: Uso de modelo local de embeddings (all-MiniLM-L6-v2), limitación de tasa por usuario y cuota máxima diaria institucional.",
         "Contingencia: Exponential Backoff con Jitter ante errores HTTP 429 y conmutación automática (failover) a proveedor secundario o modelo local (Ollama/vLLM)."),

        ("RSK-05: Inyección indirecta de prompts en documentos",
         "Mitigación: Encapsulamiento del texto de documentos en etiquetas XML pasivas (<contexto_documental>) e instrucción estricta al System Prompt de no ejecutar órdenes internas.",
         "Contingencia: Validador de salida para detectar patrones de fuga de prompts del sistema y bloqueo preventivo con alerta en auditoría."),

        ("RSK-06: Fuga de información entre roles en búsqueda RAG",
         "Mitigación: Payload Filtering en base de datos vectorial aplicando filtros booleanos de roles y usuarios antes de calcular similitud.",
         "Contingencia: Doble verificación en backend validando los permisos del documento de cada chunk antes de armar el prompt generativo.")
    ]

    for title, mit, cont in rsks:
        add_heading_2(doc, title)
        add_p(doc, mit)
        add_p(doc, cont)

    doc.save(os.path.join(OUTPUT_DIR, "06_Analisis_de_Riesgos.docx"))
    print("Doc 6 generated.")

# ==============================================================================
# DOC 7: MATRIZ DE TRAZABILIDAD
# ==============================================================================
def build_doc7():
    doc = create_base_doc()
    add_header_block(doc, "MATRIZ DE TRAZABILIDAD DE REQUISITOS (RTM)",
                     "Sistema Inteligente de Gestión y Análisis Documental | Unidades Tecnológicas de Santander (UTS)\nEstándar: Requirements Traceability Matrix Bidireccional")

    add_heading_1(doc, "1. Marco Metodológico de la Trazabilidad")
    add_p(doc, "La matriz de trazabilidad garantiza la alineación bidireccional entre los objetivos de negocio de las UTS, los requisitos funcionales (RF), las historias de usuario (HU), los casos de uso (CU), los módulos de arquitectura y las pruebas de aseguramiento de calidad (QA).")

    add_heading_1(doc, "2. Tabla Maestra de Trazabilidad Cruzada (Funcional)")
    headers = ["ID RF", "Objetivo UTS", "HU", "CU", "Capa / Módulo de Software", "Mecanismo de Verificación", "Estado"]
    rows = [
        ["RF-01", "OE-01: Seguridad", "HU-01", "CU-01", "Auth Controller / Spring Security / DB", "Pruebas Unitarias JUnit + Integration JWT", "Aprobado"],
        ["RF-02", "OE-01: Repositorio", "HU-02", "CU-05", "Folder Service / Spring Boot / DB", "Pruebas Unitarias + Test CRUD", "Aprobado"],
        ["RF-03", "OE-01: Ingesta", "HU-03", "CU-02", "Multipart Ingestion / Local Storage", "Pruebas Integración Upload (20MB/Magic)", "Aprobado"],
        ["RF-04", "OE-02: Extracción/OCR", "HU-03", "CU-02/03", "FastAPI / PyMuPDF / Tesseract OCR", "PyTest Extracción + Test E2E OCR", "Aprobado"],
        ["RF-05", "OE-02: Clasificación", "HU-04", "CU-03", "Classifier Service / Scikit / LLM", "PyTest Confusión (F1 >= 0.85)", "Aprobado"],
        ["RF-06", "OE-03: Resumen", "HU-05", "CU-03", "Summarizer Service / LLM API", "Test Inferencia + Esquema Markdown", "Aprobado"],
        ["RF-07", "OE-03: Entidades JSON", "HU-06", "CU-03", "Pydantic Schema Extractor / DB", "PyTest Validación JSON Schema", "Aprobado"],
        ["RF-08", "OE-04: RAG Q&A", "HU-07", "CU-04", "RAG Service / Vector DB / LLM", "Test Integración RAG (Latencia < 2.5s)", "Aprobado"],
        ["RF-09", "OE-04: Búsqueda", "HU-07", "CU-06", "Search Service / PostgreSQL FullText", "Pruebas Unitarias SQL + Carga Search", "Aprobado"],
        ["RF-10", "OE-05: Dashboard", "HU-08", "CU-07", "Frontend Charts / Analytics Service", "Pruebas E2E Playwright / Cypress", "Aprobado"],
        ["RF-11", "OE-05: Auditoría", "HU-08", "CU-08", "Audit Interceptor / PostgreSQL Table", "Pruebas Integración Fallas / Logs", "Aprobado"]
    ]
    add_table(doc, headers, rows, [0.6, 1.1, 0.6, 0.7, 1.8, 1.8, 0.6])

    add_heading_1(doc, "3. Matriz de Trazabilidad de Requisitos No Funcionales (RNF)")
    rnf_headers = ["Código", "Categoría", "Componente Técnico", "Técnica Implementada", "Criterio de Validación QA"]
    rnf_rows = [
        ["RNF-01", "Seguridad", "Gateway, Auth, Ingestion", "BCrypt, JWT RSA256, TLS 1.3, Magic Bytes", "Escaneo SAST/SonarQube sin vulnerabilidades críticas."],
        ["RNF-02", "Rendimiento", "Vector DB, Redis, Spring", "Índice HNSW, Redis Cache, Hilos Async", "Pruebas JMeter P95 <= 2.5s con 50 usuarios concurrentes."],
        ["RNF-03", "Disponibilidad", "Docker Compose / Health", "restart: always, /actuator/health", "Recuperación automática de contenedores en <10s."],
        ["RNF-04", "Usabilidad", "Frontend React / UI", "WCAG 2.1 AA, componentes accesibles", "Prueba SUS con usuarios finales >= 80/100."],
        ["RNF-05", "Mantenibilidad", "Backend & IA Modules", "Clean Architecture, OpenAPI 3.0", "Cobertura JaCoCo / PyTest >= 80%, Sonar grado 'A'."],
        ["RNF-06", "Escalabilidad", "Worker de IA / Celery", "Desacoplamiento con colas de tareas", "Escalamiento horizontal de workers (1 a 4 réplicas)."]
    ]
    add_table(doc, rnf_headers, rnf_rows, [0.7, 1.1, 1.5, 1.8, 2.1])

    add_heading_1(doc, "4. Aprobación de la Línea Base de Análisis")
    add_p(doc, "La presente línea base de requisitos de la fase 01_ANALISIS ha sido verificada y aprobada por el Equipo de Requisitos, Arquitectura de Software y Aseguramiento de la Calidad (QA) para proceder con la fase 02_DISENO.")

    doc.save(os.path.join(OUTPUT_DIR, "07_Matriz_de_Trazabilidad.docx"))
    print("Doc 7 generated.")

if __name__ == "__main__":
    build_doc1()
    build_doc2()
    build_doc3()
    build_doc4()
    build_doc5()
    build_doc6()
    build_doc7()
    print("All 7 DOCX files have been successfully generated!")
