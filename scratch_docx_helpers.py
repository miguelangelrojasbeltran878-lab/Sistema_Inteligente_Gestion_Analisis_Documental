import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_heading(p, font_size=14, bold=True, color_rgb=(14, 56, 122), space_before=12, space_after=6):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    for run in p.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*color_rgb)

def add_header_block(doc, title, subtitle):
    # Header title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(4)
    r1 = title_p.add_run(title)
    r1.font.name = 'Arial'
    r1.font.size = Pt(18)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(14, 56, 122)

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(14)
    r2 = sub_p.add_run(subtitle)
    r2.font.name = 'Arial'
    r2.font.size = Pt(11)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(80, 80, 80)

def add_table_data(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "0E387A")
        set_cell_margins(hdr_cells[i], 120, 120, 150, 150)
        for p in hdr_cells[i].paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F4F6F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], 100, 100, 120, 120)
            for p in row_cells[c_idx].paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = 'Arial'
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(40, 40, 40)

    # Set Widths if provided
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

def add_callout(doc, text, title="NOTA IMPORTANTE:"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    
    r_t = p.add_run(f"[{title}] ")
    r_t.font.name = 'Arial'
    r_t.font.bold = True
    r_t.font.size = Pt(9.5)
    r_t.font.color.rgb = RGBColor(14, 56, 122)
    
    r_b = p.add_run(text)
    r_b.font.name = 'Arial'
    r_b.font.size = Pt(9.5)
    r_b.font.italic = True
    r_b.font.color.rgb = RGBColor(60, 60, 60)

def add_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.bold = True
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = RGBColor(20, 20, 20)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(40, 40, 40)
    return p

print("Helper functions defined.")
